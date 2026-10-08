"""CSV loading, quality audit, reference checks and vehicle-disjoint splits."""
import csv
import hashlib
import json
from pathlib import Path

import numpy as np
import pandas as pd
from sklearn.model_selection import GroupShuffleSplit

from .feature_engineering import (ALL_COLUMNS, IDENTIFIERS, NUMERIC_COLUMNS,
                                  STRING_COLUMNS, TARGETS, invalid_numeric_mask)

ROOT = Path(__file__).resolve().parents[1]
DEFAULT_DATA = ROOT / "data/raw/du_lieu_phan_loai_gia_lap.csv"
SIMULATION_WARNING = (
    "Dữ liệu và nhãn đều giả lập theo quy tắc. Kết quả không chứng minh khả năng chẩn đoán xe thật, "
    "không dùng để kết luận màn hình gây chết máy hoặc quyết định thay/nâng máy phát. "
    "Các giá trị mô phỏng dùng làm đầu vào phải được đo hoặc ước lượng hợp lệ trước dự đoán."
)


def sha256_file(path):
    digest = hashlib.sha256()
    with Path(path).open("rb") as handle:
        for block in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


def write_json(path, content):
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    def serialize(value):
        if isinstance(value, np.generic):
            return value.item()
        if isinstance(value, (Path, np.ndarray)):
            return str(value) if isinstance(value, Path) else value.tolist()
        raise TypeError(f"Cannot serialize {type(value)}")
    path.write_text(json.dumps(content, ensure_ascii=False, indent=2, default=serialize,
                               allow_nan=False), encoding="utf-8")


def load_dataset(path, training=True):
    path = Path(path)
    with path.open(encoding="utf-8-sig", newline="") as handle:
        header = next(csv.reader(handle), [])
    if len(header) != len(set(header)):
        raise ValueError("Duplicate CSV column names are not allowed.")
    df = pd.read_csv(path, encoding="utf-8-sig")
    if df.empty:
        raise ValueError("Dataset is empty.")
    if training:
        missing = sorted(set(ALL_COLUMNS) - set(df.columns))
        if missing:
            raise ValueError(f"Training schema missing columns: {missing}")
    for c in df.select_dtypes(include=["object", "string"]).columns:
        df[c] = df[c].map(lambda v: (v.strip() or np.nan) if isinstance(v, str) else v)
    return df


def validate_training_data(df):
    missing = sorted(set(ALL_COLUMNS) - set(df.columns))
    if missing:
        raise ValueError(f"Missing training columns: {missing}")
    if df[IDENTIFIERS].isna().any().any():
        raise ValueError("Missing scenario_id, vehicle_id or profile_id.")
    if df["scenario_id"].duplicated().any() or df.duplicated().any():
        raise ValueError("Duplicate scenarios/rows: fix source explicitly, no automatic deletion.")
    if (df.groupby("vehicle_id")["profile_id"].nunique() > 1).any():
        raise ValueError("A vehicle maps to multiple profiles.")
    for c in NUMERIC_COLUMNS:
        if invalid_numeric_mask(df[c], c).any():
            raise ValueError(f"Invalid numerical values in {c}; see EDA report.")
    for target in TARGETS:
        if df[target].isna().any() or set(df[target].unique()) != {0, 1}:
            raise ValueError(f"Target {target} must contain both binary classes and no missing values.")


def reference_audit(df, path):
    if not Path(path).exists():
        return {"available": False, "path": str(path)}
    reference = load_dataset(path, training=False)
    if reference["profile_id"].isna().any() or reference["profile_id"].duplicated().any():
        raise ValueError("Reference profile_id must be unique and non-missing.")
    ref = reference.set_index("profile_id")
    report = {"available": True, "shape": list(reference.shape), "sha256": sha256_file(path),
              "unknown_profiles": sorted(set(df["profile_id"].dropna()) - set(ref.index)),
              "field_mismatches": {}, "simulated_alternator_deviations": {}}
    mapping = {"hang_xe": "hang_xe", "mau_xe": "mau_xe", "doi_xe": "nam",
               "dung_tich_dong_co_L": "dung_tich_L", "xang_do_thi_L_100km": "xang_do_thi_L_100km",
               "xang_hon_hop_L_100km": "xang_hon_hop_L_100km"}
    known = df["profile_id"].isin(ref.index)
    for column, ref_column in mapping.items():
        expected = df["profile_id"].map(ref[ref_column])
        if column in STRING_COLUMNS:
            match = df[column].eq(expected) | (df[column].isna() & expected.isna())
        else:
            match = np.isclose(pd.to_numeric(df[column], errors="coerce"),
                               pd.to_numeric(expected, errors="coerce"), equal_nan=True)
        report["field_mismatches"][column] = int((known & ~match).sum())
    expected = df["profile_id"].map(ref["may_phat_A_GIA_LAP_DUNG_CHO_SIM"])
    actual = pd.to_numeric(df["may_phat_dinh_muc_A_GIA_LAP"], errors="coerce")
    deviations = known & ~np.isclose(actual, expected, equal_nan=True)
    report["simulated_alternator_deviations"] = {
        "rows": int(deviations.sum()), "vehicles": int(df.loc[deviations, "vehicle_id"].nunique()),
        "interpretation": "Simulation scenario ratings can deviate from reference; not verified real-car ratings. No merge or automatic correction."}
    return report


def workbook_audit(path, output, df):
    if not Path(path).exists():
        return {"available": False, "path": str(path)}
    from openpyxl import load_workbook
    workbook = load_workbook(path, read_only=True, data_only=False)
    report = {"available": True, "sha256": sha256_file(path), "sheets": {}, "note": "Reference only; no rows appended to training."}
    for sheet in workbook:
        rows = list(sheet.iter_rows(values_only=True))
        formulas = [{"row": i + 1, "column": j + 1, "formula": value}
                    for i, row in enumerate(rows) for j, value in enumerate(row)
                    if isinstance(value, str) and value.startswith("=")]
        report["sheets"][sheet.title] = {"rows": len(rows), "columns": sheet.max_column,
                                        "formula_cells": formulas}
        if sheet.title == "Data_Dictionary":
            dictionary = pd.DataFrame(rows[1:], columns=rows[0]).dropna(how="all")
            dictionary.to_csv(output / "data_dictionary.csv", index=False, encoding="utf-8-sig")
        if sheet.title == "Scenarios_SIM" and rows:
            scenarios = pd.DataFrame(rows[1:], columns=rows[0]).dropna(how="all")
            common = [c for c in df.columns if c in scenarios.columns]
            report["scenario_csv_comparison"] = {"rows": len(scenarios), "common_columns": len(common),
                "additional_workbook_columns": [c for c in scenarios if c not in df], "mismatched_cells": {}}
            if len(scenarios) == len(df):
                for c in common:
                    left, right = df[c].reset_index(drop=True), scenarios[c].reset_index(drop=True)
                    if c in NUMERIC_COLUMNS:
                        equal = np.isclose(pd.to_numeric(left, errors="coerce"), pd.to_numeric(right, errors="coerce"), equal_nan=True)
                    else:
                        equal = left.eq(right) | (left.isna() & right.isna())
                    count = int((~equal).sum())
                    if count:
                        report["scenario_csv_comparison"]["mismatched_cells"][c] = count
    workbook.close()
    return report


def audit_dataset(df, output, reference_path=None, workbook_path=None, source_path=None):
    output = Path(output)
    output.mkdir(parents=True, exist_ok=True)
    missing = pd.DataFrame({"column": df.columns, "missing_count": df.isna().sum().values,
                            "missing_pct": df.isna().mean().values * 100})
    missing.to_csv(output / "missing_values.csv", index=False, encoding="utf-8-sig")
    summaries, invalid, distributions, outliers = [], {}, {}, {}
    for c in df:
        row = {"column": c, "dtype": str(df[c].dtype), "count": int(df[c].count()),
               "unique": int(df[c].nunique()), "missing_count": int(df[c].isna().sum())}
        if c in NUMERIC_COLUMNS:
            invalid[c] = int(invalid_numeric_mask(df[c], c).sum())
            numeric = pd.to_numeric(df[c], errors="coerce").replace([np.inf, -np.inf], np.nan)
            row.update({k: (float(v) if pd.notna(v) else None) for k, v in numeric.describe().items()})
            q1, q3 = numeric.quantile([0.25, 0.75]); iqr = q3 - q1
            outliers[c] = {"IQR_flagged": int(((numeric < q1 - 1.5 * iqr) | (numeric > q3 + 1.5 * iqr)).sum()),
                           "policy": "Flag only; no deletion/clipping. Zero IQR can flag legitimate binary/rare events."}
        else:
            distributions[c] = df[c].fillna("<MISSING>").value_counts().to_dict()
        summaries.append(row)
    pd.DataFrame(summaries).to_csv(output / "feature_summary.csv", index=False, encoding="utf-8-sig")
    numeric = df[[c for c in NUMERIC_COLUMNS if c in df]].apply(pd.to_numeric, errors="coerce")
    numeric.corr().to_csv(output / "numeric_correlations.csv", encoding="utf-8-sig")
    numeric.hist(figsize=(18, 16), bins=24)
    import matplotlib.pyplot as plt
    plots = output / "plots"; plots.mkdir(exist_ok=True)
    plt.tight_layout(); plt.savefig(plots / "numeric_distributions.png", dpi=110); plt.close("all")
    targets = []
    for target in TARGETS:
        for label, count in df[target].value_counts(dropna=False).items():
            targets.append({"target": target, "label": str(label), "count": int(count), "pct": float(count / len(df) * 100)})
    pd.DataFrame(targets).to_csv(output / "target_distribution.csv", index=False, encoding="utf-8-sig")
    df.groupby(TARGETS, dropna=False).size().reset_index(name="count").to_csv(output / "joint_target_distribution.csv", index=False)
    for group in ["vehicle_id", "profile_id"]:
        grouped = df.groupby(group)[TARGETS].agg(["count", "sum", "mean"])
        grouped.columns = [f"{a}_{b}" for a, b in grouped.columns]
        grouped.to_csv(output / f"targets_by_{group}.csv", encoding="utf-8-sig")
    report = {"warning": SIMULATION_WARNING, "shape": list(df.shape),
              "schema": {c: str(df[c].dtype) for c in df},
              "duplicate_rows": int(df.duplicated().sum()),
              "identifiers": {c: {"missing": int(df[c].isna().sum()), "unique": int(df[c].nunique()),
                  "repeated_rows": int(df[c].duplicated().sum()),
                  "repetition_expected": c != "scenario_id"} for c in IDENTIFIERS},
              "invalid_numeric_counts": invalid, "outliers": outliers,
              "categorical_distributions": distributions,
              "replacement_character_counts": {c: int(df[c].astype(str).str.contains("\ufffd", regex=False).sum()) for c in STRING_COLUMNS},
              "positive_rates": {c: float(pd.to_numeric(df[c], errors="coerce").mean()) for c in TARGETS},
              "cleaning_policy": "Strip surrounding whitespace and treat blank strings as missing in memory only; fail on invalid ranges/IDs/labels, preserve valid extreme scenarios; train-fold imputation."}
    if source_path:
        report["source"] = {"path": str(source_path), "sha256": sha256_file(source_path)}
    if reference_path:
        report["reference"] = reference_audit(df, reference_path)
    if workbook_path:
        report["workbook"] = workbook_audit(workbook_path, output, df)
    write_json(output / "data_quality_report.json", report)
    return report


def group_split(df, seed=2026):
    """Fixed two-stage split; inspect class presence, never optimize test metrics."""
    groups = df["vehicle_id"]
    n_groups = groups.nunique()
    n_holdout = max(1, round(n_groups * 0.15))
    if n_groups - 2 * n_holdout < 2:
        raise ValueError("Not enough vehicles for train/validation/test splitting.")
    trainval, test = next(GroupShuffleSplit(n_splits=1, test_size=n_holdout, random_state=seed).split(df, groups=groups))
    train_local, val_local = next(GroupShuffleSplit(n_splits=1, test_size=n_holdout,
        random_state=seed + 1).split(df.iloc[trainval], groups=groups.iloc[trainval]))
    parts = {"train": trainval[train_local], "validation": trainval[val_local], "test": test}
    assert_group_disjoint(df, parts)
    for name, indices in parts.items():
        for target in TARGETS:
            if df.iloc[indices][target].nunique() != 2:
                raise ValueError(f"{name}/{target} missing a class; explicitly redesign group allocation before training.")
    return parts


def assert_group_disjoint(df, parts):
    names = list(parts)
    for i, name in enumerate(names):
        for other in names[i + 1:]:
            if set(df.iloc[parts[name]]["vehicle_id"]) & set(df.iloc[parts[other]]["vehicle_id"]):
                raise ValueError(f"Vehicle overlap between {name} and {other}.")
    indices = np.concatenate(list(parts.values()))
    if len(indices) != len(df) or len(set(indices)) != len(df):
        raise ValueError("Split must cover each row exactly once.")


def split_summary(df, parts):
    return {name: {"rows": len(indices), "vehicles": int(df.iloc[indices]["vehicle_id"].nunique()),
                  "profiles": int(df.iloc[indices]["profile_id"].nunique()),
                  "positive_rates": {c: float(df.iloc[indices][c].mean()) for c in TARGETS}}
            for name, indices in parts.items()}
